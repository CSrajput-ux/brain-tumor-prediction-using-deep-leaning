import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx

doc = docx.Document('Brain_MRI_Paper_Final_Corrected_v5.docx')

print("=== TABLE 9 ===")
t9 = doc.tables[9]
for r in t9.rows:
    print([c.text.strip().replace('\n', ' ') for c in r.cells])

print("\n=== TABLE 12 (TABLE A2) ===")
t12 = doc.tables[12]
for r in t12.rows:
    print([c.text.strip().replace('\n', ' ') for c in r.cells])

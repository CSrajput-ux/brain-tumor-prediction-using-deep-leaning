import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx
import difflib

v5 = docx.Document('Brain_MRI_Paper_Final_Corrected_v5.docx')
v6 = docx.Document('Brain_MRI_Paper_Final_Corrected_v6.docx')

print(f"v5 paragraphs: {len(v5.paragraphs)}, v6 paragraphs: {len(v6.paragraphs)}")
assert len(v5.paragraphs) == len(v6.paragraphs), "Paragraph count mismatch!"

print(f"v5 tables: {len(v5.tables)}, v6 tables: {len(v6.tables)}")
assert len(v5.tables) == len(v6.tables), "Table count mismatch!"

# Check paragraph diffs
changed_paragraphs = []
for idx in range(len(v5.paragraphs)):
    t5 = v5.paragraphs[idx].text
    t6 = v6.paragraphs[idx].text
    if t5 != t6:
        changed_paragraphs.append((idx, t5, t6))

print(f"\nTotal changed paragraphs: {len(changed_paragraphs)}")
for idx, t5, t6 in changed_paragraphs:
    print(f"\n--- Paragraph P{idx} ---")
    s = difflib.ndiff([t5], [t6])
    print('\n'.join(s))

# Check table diffs
changed_tables = []
for t_idx in range(len(v5.tables)):
    tbl5 = v5.tables[t_idx]
    tbl6 = v6.tables[t_idx]
    assert len(tbl5.rows) == len(tbl6.rows), f"Table {t_idx} row count mismatch!"
    assert len(tbl5.columns) == len(tbl6.columns), f"Table {t_idx} col count mismatch!"
    for r_idx in range(len(tbl5.rows)):
        for c_idx in range(len(tbl5.rows[r_idx].cells)):
            c5 = tbl5.rows[r_idx].cells[c_idx].text.strip()
            c6 = tbl6.rows[r_idx].cells[c_idx].text.strip()
            if c5 != c6:
                changed_tables.append((t_idx, r_idx, c_idx, c5, c6))

print(f"\nTotal changed table cells: {len(changed_tables)}")
for t_idx, r_idx, c_idx, c5, c6 in changed_tables:
    print(f"Table {t_idx}, Row {r_idx}, Col {c_idx}:")
    print(f"  v5: {c5}")
    print(f"  v6: {c6}")

# Verify against logged changes:
# Logged paragraph changes: P0 (FIX 9), P6 (FIX 5), P38 (FIX 1, 7), P68 (FIX 8), P133 (FIX 8), P146 (FIX 4), P178 (FIX 5), P186 (FIX 5), P187 (FIX 2, 3)
logged_p_indices = {0, 6, 38, 68, 133, 146, 178, 186, 187}
actual_p_indices = {idx for idx, _, _ in changed_paragraphs}

unlogged_p = actual_p_indices - logged_p_indices
print(f"\nUnlogged paragraph changes: {unlogged_p}")
assert len(unlogged_p) == 0, f"Unlogged paragraph changes found: {unlogged_p}"

# Logged table changes: Table 1 (Table 2 in text), Row 0 Col 3, Row 1 Col 3, Row 2 Col 3, Row 3 Col 3
logged_tbl_cells = {(1, 0, 3), (1, 1, 3), (1, 2, 3), (1, 3, 3)}
actual_tbl_cells = {(t, r, c) for t, r, c, _, _ in changed_tables}
unlogged_tbl = actual_tbl_cells - logged_tbl_cells
print(f"Unlogged table changes: {unlogged_tbl}")
assert len(unlogged_tbl) == 0, f"Unlogged table changes found: {unlogged_tbl}"

print("\n>>> AUDIT PASSED: COUNT OF UNLOGGED CHANGES IS STRICTLY ZERO (0) <<<")

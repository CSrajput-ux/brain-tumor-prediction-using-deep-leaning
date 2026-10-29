import docx
from pathlib import Path

doc = docx.Document("paper/Brain_MRI_Paper_Final_Corrected_v6.docx")
print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

print("\n=== All Headings and Structural Paragraphs ===")
for i, p in enumerate(doc.paragraphs):
    style = getattr(p.style, "name", "")
    t = p.text.strip()
    if style.startswith("Heading") or (len(t) < 70 and t.startswith(("1 ", "2 ", "3 ", "4 ", "5 ", "6 ", "7 ", "Declarations", "References", "Appendix", "Nomenclature", "Abstract", "Keywords"))):
        print(f"P{i:03d} [{style}]: {t}")

print(f"\n=== Tables Overview ({len(doc.tables)} tables) ===")
for i, tbl in enumerate(doc.tables):
    rows = len(tbl.rows)
    cols = len(tbl.columns)
    header = [c.text.strip().replace("\n", " ") for c in tbl.rows[0].cells[:4]] if rows > 0 else []
    print(f"Table {i+1} ({rows}x{cols}): {header}")

# Check embedded images in DOCX
doc_part = doc.part
images = [rel.target_ref for rel in doc_part.rels.values() if "image" in rel.target_ref]
print(f"\nEmbedded images count in DOCX: {len(images)}")
for img in images[:10]:
    print(f"  - {img}")

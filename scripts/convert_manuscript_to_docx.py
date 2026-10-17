import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path
import re

ROOT = Path(__file__).resolve().parent.parent
PUB_DIR = ROOT / 'publication' / 'manuscript_v1'

def convert_md_to_docx():
    print("Converting manuscript.md to manuscript.docx...")
    md_path = PUB_DIR / 'manuscript.md'
    docx_path = PUB_DIR / 'manuscript.docx'
    
    doc = docx.Document()
    
    # Page setup: Standard Letter, 1-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Styles setup
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    
    lines = md_path.read_text(encoding='utf-8').splitlines()
    
    in_table = False
    table_lines = []
    
    def process_table(t_lines):
        if not t_lines:
            return
        rows = []
        for line in t_lines:
            if re.match(r'^\s*\|?\s*[-:]+[-| :]*$', line):
                continue
            cells = [c.strip() for c in line.strip().strip('|').split('|')]
            if cells and any(cells):
                rows.append(cells)
                
        if not rows:
            return
            
        n_cols = max(len(r) for r in rows)
        table = doc.add_table(rows=len(rows), cols=n_cols)
        table.style = 'Light Grid Accent 1'
        
        for r_idx, row in enumerate(rows):
            for c_idx, cell in enumerate(row):
                if c_idx < n_cols:
                    table.cell(r_idx, c_idx).text = cell
                    # Formatting
                    for p in table.cell(r_idx, c_idx).paragraphs:
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)
                        for run in p.runs:
                            run.font.name = 'Times New Roman'
                            run.font.size = Pt(9.5)
                            if r_idx == 0:
                                run.font.bold = True
        doc.add_paragraph()
        
    for line in lines:
        line_s = line.strip()
        
        # Check table
        if line_s.startswith('|'):
            in_table = True
            table_lines.append(line_s)
            continue
        elif in_table:
            in_table = False
            process_table(table_lines)
            table_lines = []
            
        if not line_s:
            continue
            
        # Headers
        if line_s.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(line_s[2:].strip())
            run.font.name = 'Times New Roman'
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x11, 0x22, 0x44)
        elif line_s.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(line_s[3:].strip())
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x22, 0x33, 0x55)
        elif line_s.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line_s[4:].strip())
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x33, 0x44, 0x66)
        elif line_s.startswith('---'):
            continue
        elif line_s.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            text = line_s[2:].strip()
            # Parse simple bolding
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)
        elif re.match(r'^\d+\.\s', line_s):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            text = re.sub(r'^\d+\.\s', '', line_s)
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            parts = re.split(r'(\*\*.*?\*\*)', line_s)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)
                    
    if in_table:
        process_table(table_lines)
        
    doc.save(docx_path)
    print(f"manuscript.docx successfully generated ({docx_path.stat().st_size} bytes).")

if __name__ == '__main__':
    convert_md_to_docx()

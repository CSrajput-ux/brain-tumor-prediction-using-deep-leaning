import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx

doc = docx.Document('Brain_MRI_Paper_Final_Corrected_v6.docx')

# 1. Abstract (P6)
p6 = doc.paragraphs[6]
print("Old P6:")
print(p6.text)

# We want:
# 1) Use 'partially confounded'
# 2) Avoid saying the same thing twice back-to-back
old_p6_target1 = "which this design cannot confirm because tumour class is confounded with imaging plane"
new_p6_target1 = "which this design cannot confirm because tumour class is partially confounded with imaging plane"

old_p6_target2 = ", as tumour class is partially confounded with imaging plane."
new_p6_target2 = "."

assert old_p6_target1 in p6.text
assert old_p6_target2 in p6.text

p6_text = p6.text.replace(old_p6_target1, new_p6_target1).replace(old_p6_target2, new_p6_target2)
print("\nNew P6:")
print(p6_text)

# 2. Keywords (P7)
p7 = doc.paragraphs[7]
print("\nOld P7:")
print(p7.text)
assert "Artifact disentanglement;" in p7.text
p7_text = p7.text.replace("Artifact disentanglement;", "Seam artefacts;")
print("New P7:")
print(p7_text)

# 3. §3.1 Plane (P64)
p64 = doc.paragraphs[64]
print("\nOld P64:")
print(p64.text)
old_p64_target = "Clinical MRI protocols routinely acquire pituitary tumours in coronal and sagittal planes to evaluate the sella turcica, pituitary stalk, and optic chiasm, whereas supratentorial gliomas are predominantly imaged in axial planes."
new_p64_target = "Clinical MRI protocols typically acquire pituitary tumours in coronal and sagittal planes to evaluate the sella turcica, pituitary stalk, and optic chiasm, whereas supratentorial gliomas are typically imaged in axial planes."
assert old_p64_target in p64.text
p64_text = p64.text.replace(old_p64_target, new_p64_target)
print("New P64:")
print(p64_text)

# 4. §4.5 2nd paragraph (P147)
p147 = doc.paragraphs[147]
print("\nOld P147:")
print(p147.text)
old_p147_target = "And because every intervention pushes predictions towards one class, the networks do not fail symmetrically;"
new_p147_target = "And because most interventions push predictions mainly towards pituitary tumour, the networks do not fail symmetrically;"
assert old_p147_target in p147.text
p147_text = p147.text.replace(old_p147_target, new_p147_target)
print("New P147:")
print(p147_text)

# 5. §5.1 (ii) (P178)
p178 = doc.paragraphs[178]
print("\nOld P178:")
print(p178.text)
old_p178_target = "reflects training and fold-level variance rather than near-duplicate grouping stringency"
new_p178_target = "most likely reflects training and fold-level variance rather than near-duplicate grouping stringency"
assert old_p178_target in p178.text
p178_text = p178.text.replace(old_p178_target, new_p178_target)
print("New P178:")
print(p178_text)

# 6. §5.4 Plane (P186)
p186 = doc.paragraphs[186]
print("\nOld P186:")
print(p186.text)
old_p186_target = "Slices with pituitary tumours are predominantly acquired in coronal and sagittal orientations to capture sellar anatomy, whereas gliomas are primarily imaged axially."
new_p186_target = "Slices with pituitary tumours are typically acquired in coronal and sagittal orientations to capture sellar anatomy, whereas gliomas are typically imaged axially."
assert old_p186_target in p186.text
p186_text = p186.text.replace(old_p186_target, new_p186_target)
print("New P186:")
print(p186_text)

import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx

doc = docx.Document('Brain_MRI_Paper_Final_Corrected_v5.docx')

paragraphs_to_check = [0, 6, 7, 35, 38, 64, 68, 133, 146, 174, 178, 186, 187, 192, 209, 215]

for p_idx in paragraphs_to_check:
    print(f"\n==================== P{p_idx} ====================")
    print(doc.paragraphs[p_idx].text)

print("\n==================== HEADERS & FOOTERS ====================")
for s_idx, section in enumerate(doc.sections):
    print(f"Section {s_idx} Header: {[p.text for p in section.header.paragraphs if p.text.strip()]}")
    print(f"Section {s_idx} Footer: {[p.text for p in section.footer.paragraphs if p.text.strip()]}")

print("\n==================== TABLE 1 (TABLE 2 IN TEXT) ====================")
t2 = doc.tables[1]
for r_idx, row in enumerate(t2.rows):
    print(f"Row {r_idx}: {[c.text.strip() for c in row.cells]}")

import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx

doc = docx.Document('Brain_MRI_Paper_Final_Corrected_v5.docx')

print("=== CHECKING SECTION 4.8 & TABLE 8 ===")
for idx, p in enumerate(doc.paragraphs):
    if "4.8" in p.text or "Table 8" in p.text or "Table 7" in p.text:
        print(f"P{idx}: {p.text}")
        for j in range(idx+1, min(idx+5, len(doc.paragraphs))):
            print(f"   P{j}: {doc.paragraphs[j].text}")

print("\n=== TABLE 8 ===")
t8 = doc.tables[8]
for r in t8.rows:
    print([c.text.strip() for c in r.cells])

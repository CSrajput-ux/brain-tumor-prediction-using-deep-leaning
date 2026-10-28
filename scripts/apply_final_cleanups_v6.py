import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx
import shutil
import hashlib
import json
from pathlib import Path

BASE_DIR = Path(r"c:\Users\jkgga\Downloads\ChengResearch")
V6_PATH = BASE_DIR / "Brain_MRI_Paper_Final_Corrected_v6.docx"
BACKUP_PATH = BASE_DIR / "Brain_MRI_Paper_Final_Corrected_v6_pre_cleanup.docx"

shutil.copyfile(V6_PATH, BACKUP_PATH)
print(f"Created backup at {BACKUP_PATH}")

doc = docx.Document(V6_PATH)

# 1. Abstract (P6)
p6 = doc.paragraphs[6]
old_p6_target1 = "which this design cannot confirm because tumour class is confounded with imaging plane"
new_p6_target1 = "which this design cannot confirm because tumour class is partially confounded with imaging plane"
old_p6_target2 = ", as tumour class is partially confounded with imaging plane."
new_p6_target2 = "."

assert old_p6_target1 in p6.text, "Target 1 not in P6"
assert old_p6_target2 in p6.text, "Target 2 not in P6"
p6.text = p6.text.replace(old_p6_target1, new_p6_target1).replace(old_p6_target2, new_p6_target2)
print("1. Abstract (P6) updated: unified to 'partially confounded' and eliminated repetition.")

# 2. Keywords (P7)
p7 = doc.paragraphs[7]
assert "Artifact disentanglement;" in p7.text, "Artifact disentanglement not in P7"
p7.text = p7.text.replace("Artifact disentanglement;", "Seam artefacts;")
print("2. Keywords (P7) updated: replaced 'Artifact disentanglement;' with 'Seam artefacts;' to match title.")

# 3. §3.1 Plane claim (P64)
p64 = doc.paragraphs[64]
old_p64_target = "Clinical MRI protocols routinely acquire pituitary tumours in coronal and sagittal planes to evaluate the sella turcica, pituitary stalk, and optic chiasm, whereas supratentorial gliomas are predominantly imaged in axial planes."
new_p64_target = "Clinical MRI protocols typically acquire pituitary tumours in coronal and sagittal planes to evaluate the sella turcica, pituitary stalk, and optic chiasm, whereas supratentorial gliomas are typically imaged in axial planes."
assert old_p64_target in p64.text, "Plane claim not in P64"
p64.text = p64.text.replace(old_p64_target, new_p64_target)
print("3. §3.1 (P64) updated: softened plane claims to 'typically'.")

# 4. §4.5 2nd paragraph (P147)
p147 = doc.paragraphs[147]
old_p147_target = "And because every intervention pushes predictions towards one class, the networks do not fail symmetrically;"
new_p147_target = "And because most interventions push predictions mainly towards pituitary tumour, the networks do not fail symmetrically;"
assert old_p147_target in p147.text, "Target not in P147"
p147.text = p147.text.replace(old_p147_target, new_p147_target)
print("4. §4.5 (P147) updated: changed to 'most interventions push predictions mainly towards pituitary tumour'.")

# 5. §5.1 (ii) (P178)
p178 = doc.paragraphs[178]
old_p178_target = "reflects training and fold-level variance rather than near-duplicate grouping stringency"
new_p178_target = "most likely reflects training and fold-level variance rather than near-duplicate grouping stringency"
assert old_p178_target in p178.text, "Target not in P178"
p178.text = p178.text.replace(old_p178_target, new_p178_target)
print("5. §5.1 (P178) updated: added 'most likely' to hedge variance explanation.")

# 6. §5.4 Plane claim (P186)
p186 = doc.paragraphs[186]
old_p186_target = "Slices with pituitary tumours are predominantly acquired in coronal and sagittal orientations to capture sellar anatomy, whereas gliomas are primarily imaged axially."
new_p186_target = "Slices with pituitary tumours are typically acquired in coronal and sagittal orientations to capture sellar anatomy, whereas gliomas are typically imaged axially."
assert old_p186_target in p186.text, "Plane claim not in P186"
p186.text = p186.text.replace(old_p186_target, new_p186_target)
print("6. §5.4 (P186) updated: softened plane claims to 'typically'.")

# 7. §7 Conclusion (P215)
p215 = doc.paragraphs[215]
if "tumour class is potentially confounded with imaging plane." in p215.text:
    p215.text = p215.text.replace("tumour class is potentially confounded with imaging plane.", "tumour class is partially confounded with imaging plane.")
    print("7. §7 (P215) updated: aligned to 'partially confounded'.")

# Save updated v6
doc.save(str(V6_PATH))
print(f"Successfully saved clean v6 to {V6_PATH}!")

# Sync to publication/manuscript_v1/manuscript.docx
pub_m = BASE_DIR / "publication" / "manuscript_v1" / "manuscript.docx"
shutil.copyfile(V6_PATH, pub_m)

# Update checksums.json
h = hashlib.sha256()
with open(pub_m, 'rb') as f:
    while chunk := f.read(8192):
        h.update(chunk)
sha = h.hexdigest()
size = pub_m.stat().st_size

with open(BASE_DIR / 'publication' / 'manuscript_v1' / 'checksums.json', 'r') as f:
    ck = json.load(f)

ck['manuscript.docx'] = {'sha256': sha, 'size_bytes': size}

with open(BASE_DIR / 'publication' / 'manuscript_v1' / 'checksums.json', 'w') as f:
    json.dump(ck, f, indent=2)

print(f"Synced publication manuscript.docx: {sha[:12]} ({size} bytes).")

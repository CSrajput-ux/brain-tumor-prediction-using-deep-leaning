import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx
from docx.shared import Pt
import shutil
from pathlib import Path

BASE_DIR = Path(r"c:\Users\jkgga\Downloads\ChengResearch")
V5_PATH = BASE_DIR / "Brain_MRI_Paper_Final_Corrected_v5.docx"
V6_PATH = BASE_DIR / "Brain_MRI_Paper_Final_Corrected_v6.docx"

doc = docx.Document(V5_PATH)

# ==============================================================================
# FIX 9: Title (P0)
# ==============================================================================
p0 = doc.paragraphs[0]
old_p0 = p0.text
p0.text = "Seam Artefacts and Spatial-Context Sensitivity in Brain MRI Tumour Classification: A Multi-Control Spatial Intervention Study"
print("FIX 9 applied: Title updated.")

# ==============================================================================
# FIX 5 (Part a, c): Abstract (P6)
# ==============================================================================
p6 = doc.paragraphs[6]
old_p6 = p6.text

# Replace "consistent with loss of orientation and cranial-contour priors"
target_p6_cue = "consistent with loss of orientation and cranial-contour priors"
repl_p6_cue = "a pattern that could reflect loss of orientation cues, which this design cannot confirm because tumour class is confounded with imaging plane"
assert target_p6_cue in p6.text, f"Target not in P6: {target_p6_cue}"
p6_text_step1 = p6.text.replace(target_p6_cue, repl_p6_cue)

# Replace "These results demonstrate the co-presence of both boundary-artifact sensitivity and distant-content sensitivity"
target_p6_demo = "These results demonstrate the co-presence of both boundary-artifact sensitivity and distant-content sensitivity"
repl_p6_demo = "These results are consistent with the presence of both boundary-artifact sensitivity and distant-content sensitivity"
assert target_p6_demo in p6_text_step1, f"Target not in P6: {target_p6_demo}"
p6_text_step2 = p6_text_step1.replace(target_p6_demo, repl_p6_demo)

p6.text = p6_text_step2
print("FIX 5 applied to Abstract (P6).")

# ==============================================================================
# FIX 1 & FIX 7: Hypotheses Paragraph (P38)
# ==============================================================================
p38 = doc.paragraphs[38]
old_p38 = p38.text

# FIX 7: Add pre-registration caveat after first sentence
target_p38_lead = "The hypotheses were formulated in the phased research protocol before model execution:"
repl_p38_lead = "The hypotheses were formulated in the phased research protocol before model execution. They were not deposited on an immutable external registry, and repository timestamps cannot prove the order of configuration and predictions. They were:"
assert target_p38_lead in p38.text, f"Target not in P38: {target_p38_lead}"
p38_text_step1 = p38.text.replace(target_p38_lead, repl_p38_lead)

# FIX 1: Replace clinical margin / SE comparison with practical-significance threshold
target_p38_thresh = "a pre-specified clinical margin of 0.150 Macro-F1 (approximately 2.5 times the fold-level standard error, SE ≈ 0.015)"
repl_p38_thresh = "a pre-specified practical-significance threshold of 0.150 Macro-F1 (a project convention, not a clinically derived margin)"
assert target_p38_thresh in p38_text_step1, f"Target not in P38: {target_p38_thresh}"
p38_text_step2 = p38_text_step1.replace(target_p38_thresh, repl_p38_thresh)

p38.text = p38_text_step2
print("FIX 1 & FIX 7 applied to Hypotheses (P38).")

# ==============================================================================
# FIX 6: Table 1 in docx (Table 2 in text) - Imaging Plane Column
# ==============================================================================
t2 = doc.tables[1]

# Header (Row 0, Cell 3)
t2.rows[0].cells[3].text = "Typical anatomical location and commonly used imaging planes (not verified in this dataset)"

# Meningioma (Row 1, Cell 3)
t2.rows[1].cells[3].text = "Dural-based, extra-axial (parasagittal, convexity); typically axial & coronal views [AUTHOR ACTION: cite source or remove plane claim]"

# Glioma (Row 2, Cell 3)
t2.rows[2].cells[3].text = "Intra-axial cerebral/cerebellar parenchyma; typically axial views [AUTHOR ACTION: cite source or remove plane claim]"

# Pituitary tumour (Row 3, Cell 3)
t2.rows[3].cells[3].text = "Sella turcica & skull base; typically coronal & sagittal views [AUTHOR ACTION: cite source or remove plane claim]"

# Preserve font size
for row in t2.rows:
    for c in row.cells:
        for p in c.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(8.5)

print("FIX 6 applied to Table 2 (Table 1 in docx).")

# ==============================================================================
# FIX 8 (Part b): Section 3.2 Evaluation scenarios (P68)
# ==============================================================================
p68 = doc.paragraphs[68]
old_p68 = p68.text
target_p68 = "reflect fold-level checkpoint training variance rather than grouping stringency"
repl_p68 = "most likely reflect fold-level checkpoint training variance rather than grouping stringency"
assert target_p68 in p68.text, f"Target not in P68: {target_p68}"
p68.text = p68.text.replace(target_p68, repl_p68)
print("FIX 8 applied to Section 3.2 (P68).")

# ==============================================================================
# FIX 8 (Part a): Section 4.3 (P133)
# ==============================================================================
p133 = doc.paragraphs[133]
old_p133 = p133.text
target_p133 = "and it is smaller in the scenarios with stricter near-duplicate grouping."
repl_p133 = "and it differs between scenarios; because S1 merges only two clusters relative to S0, this difference most likely reflects checkpoint-level variance rather than grouping stringency."
assert target_p133 in p133.text, f"Target not in P133: {target_p133}"
p133.text = p133.text.replace(target_p133, repl_p133)
print("FIX 8 applied to Section 4.3 (P133).")

# ==============================================================================
# FIX 4: Section 4.5 Class-transition statement (P146)
# ==============================================================================
p146 = doc.paragraphs[146]
old_p146 = p146.text
target_p146 = "Conversely, pituitary slices labelled glioma remain near zero across all conditions (0.0% under C2 and 0.9% under C3), demonstrating that distant scrambling systematically pushes predictions toward the pituitary prior."
repl_p146 = "The largest single misclassification transition is glioma to pituitary tumour. However, glioma misclassification as meningioma is also substantial under C1–C3 in S0: 374/1,426 (26.2%) for C1, 397/1,426 (27.8%) for C2, and 360/1,426 (25.2%) for C3. Conversely, pituitary slices misclassified as glioma remain low across most conditions (e.g. 0/930 under C2 and 8/930 under C3), though elevated under C1 in S0 (71/930, 7.6%)."
assert target_p146 in p146.text, f"Target not in P146: {target_p146}"
p146.text = p146.text.replace(target_p146, repl_p146)
print("FIX 4 applied to Section 4.5 (P146).")

# ==============================================================================
# FIX 5 (Part b, d, e): Section 5.1 (P178)
# ==============================================================================
p178 = doc.paragraphs[178]
old_p178 = p178.text

target_p178_estab = "establishing that boundary discontinuities alone suffice to degrade frozen checkpoints."
repl_p178_estab = "indicating that seam-like discontinuities alone can degrade these frozen checkpoints."
assert target_p178_estab in p178.text, f"Target not in P178: {target_p178_estab}"
p178_step1 = p178.text.replace(target_p178_estab, repl_p178_estab)

target_p178_demo = "Together, these observations demonstrate the co-presence of both boundary-artifact sensitivity and distant-content sensitivity, while demonstrating that complete mathematical apportionment requires a full 2×2 factorial design with energy-matched controls."
repl_p178_demo = "Together, these observations are consistent with the presence of both boundary-artifact sensitivity and distant-content sensitivity, while indicating that apportionment would require a full 2×2 factorial design with energy-matched controls."
assert target_p178_demo in p178_step1, f"Target not in P178: {target_p178_demo}"
p178_step2 = p178_step1.replace(target_p178_demo, repl_p178_demo)

p178.text = p178_step2
print("FIX 5 applied to Section 5.1 (P178).")

# ==============================================================================
# FIX 5 (Part d): Section 5.4 Plane cues (P186)
# ==============================================================================
p186 = doc.paragraphs[186]
old_p186 = p186.text
target_p186 = "is exactly what the loss of plane-specific anatomical priors predicts."
repl_p186 = "is compatible with the loss of plane-specific anatomical cues, which remains untested."
assert target_p186 in p186.text, f"Target not in P186: {target_p186}"
p186.text = p186.text.replace(target_p186, repl_p186)
print("FIX 5 applied to Section 5.4 (P186).")

# ==============================================================================
# FIX 2 & FIX 3: Section 5.4 (P187)
# ==============================================================================
p187 = doc.paragraphs[187]
old_p187 = p187.text

# FIX 3: Input HF energy values
target_p187_hf = "Bilinear resizing to 224×224 partially attenuates this thin boundary feature (reducing input HF energy to 2167 and 3073)."
repl_p187_hf = "Bilinear resizing to 224×224 attenuates this thin boundary feature, bringing input HF energy (1558 for C4a and 2107 for C4b [AUTHOR CHECK: confirm against archived diagnostics output]) below that of C3 (2410)."
assert target_p187_hf in p187.text, f"Target not in P187: {target_p187_hf}"
p187_step1 = p187.text.replace(target_p187_hf, repl_p187_hf)

# FIX 2: Unsupported dose claim
target_p187_dose = "Finally, perturbation dose varies systematically with tumour size: compact pituitary tumours leave ~91% of the background perturbed, whereas large gliomas protect larger contiguous margins (~88% perturbed), contributing to class-asymmetric sensitivity."
repl_p187_dose = "Whether perturbation dose differs by tumour class (the protected-block size varies with tumour size) was not analysed."
assert target_p187_dose in p187_step1, f"Target not in P187: {target_p187_dose}"
p187_step2 = p187_step1.replace(target_p187_dose, repl_p187_dose)

p187.text = p187_step2
print("FIX 2 & FIX 3 applied to Section 5.4 (P187).")

# Save v6
doc.save(str(V6_PATH))
print(f"Successfully saved {V6_PATH}!")

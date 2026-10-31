import docx, re, sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document("brain _tumar_research.docx")
print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

print("\n--- P00 to P05 ---")
for i in range(6):
    print(f"P{i:02d}: {doc.paragraphs[i].text}")

print("\n--- Tables Check ---")
for i, t in enumerate(doc.tables):
    rows = len(t.rows)
    cols = len(t.columns)
    header = [c.text.strip().replace('\n', ' ') for c in t.rows[0].cells[:4]]
    print(f"Table {i+1}: {rows} rows x {cols} cols | Header: {header}")

print("\n--- Table 3 (Full Plane Stratification) ---")
for r in doc.tables[2].rows:
    print([c.text.strip().replace('\n', ' ') for c in r.cells])

print("\n--- Table 4 (Full Class Heterogeneity) ---")
for r in doc.tables[3].rows:
    print([c.text.strip().replace('\n', ' ') for c in r.cells])

print("\n--- Declarations & Placeholders Check ---")
for i in range(170, 180):
    print(f"P{i:03d}: {doc.paragraphs[i].text[:90]}...")

# Check for bracketed placeholders
full_text = "\n".join([p.text for p in doc.paragraphs])
placeholders = re.findall(r"\[(?:to be inserted|todo|insert|tbd|\?\?\?)[^\]]*\]", full_text, flags=re.I)
print(f"\nRemaining bracketed placeholders: {placeholders}")


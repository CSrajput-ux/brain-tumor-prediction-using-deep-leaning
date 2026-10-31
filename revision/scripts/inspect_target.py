import docx, re, sys
from pathlib import Path
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document("brain _tumar_research.docx")
print(f"Total paragraphs: {len(doc.paragraphs)}, Total tables: {len(doc.tables)}")

img_paras = []
for i, p in enumerate(doc.paragraphs):
    if "graphic" in p._p.xml:
        rids = re.findall(r'r:embed="([^"]+)"', p._p.xml)
        caption = doc.paragraphs[i+1].text[:80] if i+1 < len(doc.paragraphs) else ""
        print(f"P{i:03d}: rIds={rids} | Caption: {caption}")
        img_paras.append(i)

print("Total image paragraphs found:", len(img_paras))
